from __future__ import annotations

import hashlib
import re
import io

from .models import Chunk, Document


def parse_bytes(name: str, data: bytes) -> Document:
    suffix = name.lower().rsplit(".", 1)[-1] if "." in name else "txt"
    if suffix in {"txt", "md"}:
        text = data.decode("utf-8", errors="replace")
    elif suffix == "pdf":
        from pypdf import PdfReader
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    elif suffix == "docx":
        from docx import Document as DocxDocument
        text = "\n".join(paragraph.text for paragraph in DocxDocument(io.BytesIO(data)).paragraphs)
    else:
        raise ValueError("Unsupported file type. Use TXT, Markdown, PDF, or DOCX.")
    normalized = re.sub(r"\s+", " ", text).strip()
    digest = hashlib.sha1(data).hexdigest()[:12]
    return Document(id=f"doc-{digest}", name=name, text=normalized)


def chunk_document(document: Document, size: int = 700, overlap: int = 100) -> list[Chunk]:
    if size <= overlap:
        raise ValueError("size must be greater than overlap")
    chunks: list[Chunk] = []
    start = 0
    index = 0
    while start < len(document.text):
        end = min(len(document.text), start + size)
        chunks.append(Chunk(id=f"{document.id}:{index}", document_id=document.id, text=document.text[start:end], start=start, end=end))
        if end == len(document.text):
            break
        start = end - overlap
        index += 1
    return chunks
